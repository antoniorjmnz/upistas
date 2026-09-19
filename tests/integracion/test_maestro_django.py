"""El maestro propio: nuestras tablas como fuente de proveedores y pedidos, y la carga del Excel."""
from dataclasses import replace
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from django.core.management import call_command

from upistas.adaptadores.fuentes.excel import MaestroExcel
from upistas.adaptadores.persistencia.django_maestro import MaestroDjango
from upistas.config import settings
from web.panel.models import Pedido, Proveedor

pytestmark = pytest.mark.django_db

REAL = settings.caja_dir / "FINAL_v7_DEFINITIVO_ahorasi.xlsx"


@pytest.fixture
def maestro():
    """Dos proveedores y tres pedidos, uno marcado para que lo mire Alberto."""
    levante = Proveedor.objects.create(
        codigo="P001", nombre="Suministros Levante S.L.", nif="b46102331",
        iban="ES21 0049 1500 0512 3456 7890", ciudad="Valencia", condiciones_dias=60,
    )
    cieza = Proveedor.objects.create(codigo="P003", nombre="Ofimática Cieza S.L.", nif="B30455812",
                                     iban="ES6001825322180201588391")
    Pedido.objects.create(numero="PO-2026-0001", proveedor=levante, importe=Decimal("2490.00"), fecha=date(2026, 1, 8))
    Pedido.objects.create(numero="PO-2026-0007", proveedor=levante, importe=Decimal("880.55"), revisar=True)
    Pedido.objects.create(numero="PO-2026-0301", proveedor=cieza, importe=Decimal("1210.00"))
    return levante


def test_guarda_el_nif_y_el_iban_como_los_comparan_las_reglas(maestro):
    assert maestro.nif == "B46102331"
    assert maestro.iban == "ES2100491500051234567890"


def test_da_los_proveedores_y_los_pedidos_del_dominio(maestro):
    fuente = MaestroDjango()
    proveedores = {p.id: p for p in fuente.proveedores()}
    assert set(proveedores) == {"P001", "P003"}
    assert proveedores["P001"].nif == "B46102331" and proveedores["P001"].condiciones_dias == 60
    assert proveedores["P003"].ciudad == "" and proveedores["P003"].condiciones_dias is None

    pedidos = {p.id: p for p in fuente.pedidos()}
    assert set(pedidos) == {"PO-2026-0001", "PO-2026-0007", "PO-2026-0301"}
    assert pedidos["PO-2026-0001"].importe == Decimal("2490.00")
    assert pedidos["PO-2026-0001"].fecha == date(2026, 1, 8)
    assert pedidos["PO-2026-0001"].proveedor_id == "P001"
    assert pedidos["PO-2026-0301"].nif == "B30455812"  # el NIF ya no se repite: lo pone el proveedor
    assert pedidos["PO-2026-0007"].fecha is None


def test_lo_que_alberto_quiere_mirar(maestro):
    assert MaestroDjango().marcados_para_revisar() == frozenset({"PO-2026-0007"})


def test_la_version_no_cambia_si_no_cambia_nada(maestro):
    assert MaestroDjango().version == MaestroDjango().version
    assert len(MaestroDjango().version) == 12


def test_la_version_cambia_en_cuanto_alberto_corrige_un_dato(maestro):
    antes = MaestroDjango().version

    maestro.iban = "ES9121000418450200051332"
    maestro.save()
    despues_de_la_cuenta = MaestroDjango().version
    assert despues_de_la_cuenta != antes

    pedido = Pedido.objects.get(numero="PO-2026-0301")
    pedido.revisar = True
    pedido.save()
    assert MaestroDjango().version != despues_de_la_cuenta


def test_la_version_de_una_instancia_es_la_foto_con_la_que_empezo(maestro):
    fuente = MaestroDjango()
    version = fuente.version
    Pedido.objects.filter(numero="PO-2026-0301").update(importe=Decimal("2000.00"))
    assert fuente.version == version and MaestroDjango().version != version


def test_el_contenedor_prefiere_el_maestro_de_la_web(maestro):
    from upistas.infra import contenedor

    original = contenedor.settings
    contenedor.configurar(replace(original, excel_path=None))
    try:
        assert isinstance(contenedor.maestro(), MaestroDjango)
    finally:
        contenedor.configurar(original)


def test_sin_proveedores_dados_de_alta_el_maestro_no_vive_aqui():
    from upistas.adaptadores.fuentes.memoria import MaestroEnMemoria
    from upistas.infra import contenedor

    original = contenedor.settings
    contenedor.configurar(replace(original, excel_path=None))
    try:
        assert not MaestroDjango.hay_datos()
        assert isinstance(contenedor.maestro(), MaestroEnMemoria)  # el conftest deja el Excel sin rutas
    finally:
        contenedor.configurar(original)


# --- la carga del Excel de La Caja ------------------------------------------------------------


@pytest.fixture
def excel_pequeno(tmp_path):
    """Un Excel con las hojas que espera `MaestroExcel`, en pequeño."""
    libro = openpyxl.Workbook()
    proveedores = libro.active
    proveedores.title = "Proveedores"
    proveedores.append(["ID", "Razon Social", "NIF", "IBAN", "Ciudad", "Condiciones"])
    proveedores.append(["P001", "Suministros Levante S.L.", "B46102331", "ES21 0049 1500 0512 3456 7890", "Valencia", "60 dias"])
    proveedores.append(["P003", "Ofimática Cieza S.L.  ", "B30455812", "es60 0182 5322 1802 0158 8391", "Murcia", "30 dias"])
    pedidos = libro.create_sheet("Pedidos_2026")
    pedidos.append(["Pedido", "ProveedorID", "NIF", "Importe_Total", "Estado", "Fecha_Pedido"])
    pedidos.append(["PO-2026-0001", "P001", "B46102331", 2490.00, "ABIERTO", "2026-01-08"])
    pedidos.append(["PO-2026-0007", "P001", "B46102331", 880.55, "ABIERTO", "2026-01-20"])
    pedidos.append(["PO-2026-0301", "P003", "B30455812", 1210.00, "ABIERTO", "2026-03-01"])
    pedidos.append(["PO-2026-0999", "P404", None, 100.00, "ABIERTO", None])  # proveedor que no existe
    revisar = libro.create_sheet("pendiente_revisar")
    revisar.append(["PO-2026-0007"])
    ruta = tmp_path / "maestro.xlsx"
    libro.save(ruta)
    libro.close()
    return ruta


def test_importar_maestro_carga_el_excel_en_nuestras_tablas(excel_pequeno, capsys):
    call_command("importar_maestro", excel=str(excel_pequeno))

    assert Proveedor.objects.count() == 2 and Pedido.objects.count() == 3
    levante = Proveedor.objects.get(codigo="P001")
    assert levante.nombre == "Suministros Levante S.L." and levante.condiciones_dias == 60
    assert Proveedor.objects.get(codigo="P003").iban == "ES6001825322180201588391"
    assert Pedido.objects.get(numero="PO-2026-0001").importe == Decimal("2490.00")
    assert Pedido.objects.get(numero="PO-2026-0001").proveedor == levante
    assert set(MaestroDjango().marcados_para_revisar()) == {"PO-2026-0007"}

    salida = capsys.readouterr().out
    assert "Proveedores: 2 (2 nuevos)" in salida and "Pedidos: 3 (3 nuevos)" in salida
    assert "Marcados para revisar: 1" in salida and "PO-2026-0999" in salida


def test_importar_maestro_dos_veces_no_duplica_y_actualiza(excel_pequeno, capsys):
    call_command("importar_maestro", excel=str(excel_pequeno))
    libro = openpyxl.load_workbook(excel_pequeno)
    libro["Proveedores"]["B2"] = "Suministros Levante S.A."
    libro["Pedidos_2026"]["D2"] = 2500.00
    libro.save(excel_pequeno)
    capsys.readouterr()

    call_command("importar_maestro", excel=str(excel_pequeno))

    assert Proveedor.objects.count() == 2 and Pedido.objects.count() == 3
    assert Proveedor.objects.get(codigo="P001").nombre == "Suministros Levante S.A."
    assert Pedido.objects.get(numero="PO-2026-0001").importe == Decimal("2500.00")
    assert "Proveedores: 2 (0 nuevos)" in capsys.readouterr().out


def test_importar_maestro_avisa_si_no_encuentra_el_excel(tmp_path):
    from django.core.management.base import CommandError

    with pytest.raises(CommandError, match="No encuentro el Excel"):
        call_command("importar_maestro", excel=str(tmp_path / "no_esta.xlsx"))


@pytest.mark.skipif(not REAL.exists(), reason="La Caja no está clonada")
def test_importar_el_maestro_de_verdad_de_la_caja(capsys):
    call_command("importar_maestro", excel=str(REAL))

    excel = MaestroExcel(REAL)
    assert Proveedor.objects.count() == 11
    assert Pedido.objects.count() == sum(1 for p in excel.pedidos() if p.proveedor_id)
    assert MaestroDjango().marcados_para_revisar() == excel.marcados_para_revisar()
    assert Pedido.objects.get(numero="PO-2026-0497").importe == Decimal("84700.00")
    assert "Proveedores: 11 (11 nuevos)" in capsys.readouterr().out
