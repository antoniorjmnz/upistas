from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from upistas.adaptadores.fuentes.excel import MaestroExcel
from upistas.config import settings

REAL = settings.caja_dir / "FINAL_v7_DEFINITIVO_ahorasi.xlsx"


@pytest.fixture
def excel_caotico(tmp_path):
    """Un Excel pequeño con las mismas trampas que el de Alberto."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Proveedores"
    ws.append(["ID", "Razon Social", "NIF", "IBAN", "Ciudad", "Condiciones"])
    ws.append(["P001", "Suministros Levante S.L.", "B46102331", "ES21 0049 1500 0512 3456 7890", "Valencia", "60 dias"])
    ws.append(["P003", "Ofimática Cieza S.L.  ", "B30455812", "es60 0182 5322 1802 0158 8391", "Murcia", "60 dias"])
    ws.append(["P007", "Papelería Ruzafa S.C.", "J40112358", "ES55 3159 0012 3487 6512 3407", "Valencia", "30 dias"])
    ws.append(["P007", "Papelería Ruzafa S.C.", "J40112358", "ES55 3159 0012 3487 6512 3407", "Valencia", "30 dias"])
    ws.append([None, None, None, None, None, None])
    p = wb.create_sheet("Pedidos_2026")
    p.append(["Pedido", "ProveedorID", "NIF", "Importe_Total", "Estado", "Fecha_Pedido"])
    p.append(["PO-2026-0001", "P003", "B30455812", 9221.75, "ABIERTO", "2026-01-31"])
    p.append(["PO-2026-0060", "P010", "B98455101", 8413, "ABIERTO", "2026-02-01"])
    p.append(["PO-2026-0546", "P007", None, 2738.78, "ABIERTO", "2026-04-23"])
    p.append(["PO-2026-0001", "P003", "B30455812", 9221.75, "ABIERTO", "2026-01-31"])
    p.append(["(archivo parcial)", None, None, None, None, None])
    old = wb.create_sheet("Pedidos_2025_OLD")
    old.append(["Pedido", "Importe"])
    old.append(["PO-2025-0812", 4100.5])
    old.append(["(archivo parcial, resto en backup_marzo??)", None])
    for basura in ("NO_TOCAR", "MACROS_ROTAS", "Hoja1", "Sheet3"):
        wb.create_sheet(basura).append(["#REF!"])
    n = wb.create_sheet("Norma_Pagos_v3")
    n.append(["NORMA DE PAGOS A PROVEEDORES (v3, vigente)"])
    n.append(["1. Pagar solo si el NIF esta en el maestro."])
    r = wb.create_sheet("pendiente_revisar")
    r.append(["PO-2026-0007"])
    r.append(["PO-2026-0141"])
    r.append(["mirar cuando haya hueco"])
    ruta = tmp_path / "caotico.xlsx"
    wb.save(ruta)
    return MaestroExcel(ruta)


def test_proveedores_limpios_y_sin_duplicados(excel_caotico):
    provs = {p.id: p for p in excel_caotico.proveedores()}
    assert set(provs) == {"P001", "P003", "P007"}
    assert provs["P003"].nombre == "Ofimática Cieza S.L."
    assert provs["P003"].iban == "ES6001825322180201588391"
    assert provs["P001"].condiciones_dias == 60
    assert any("P007 repetido" in a for a in excel_caotico.avisos)


def test_pedidos_con_importes_exactos_fechas_y_sin_nif(excel_caotico):
    pedidos = {p.id: p for p in excel_caotico.pedidos()}
    assert pedidos["PO-2026-0001"].importe == Decimal("9221.75")
    assert pedidos["PO-2026-0060"].importe == Decimal("8413.00")
    assert pedidos["PO-2026-0001"].fecha == date(2026, 1, 31)
    assert pedidos["PO-2026-0546"].nif == ""
    assert pedidos["PO-2025-0812"].estado == "ARCHIVADO_2025"
    assert "(archivo parcial)" not in pedidos
    assert any("PO-2026-0001 repetido" in a for a in excel_caotico.avisos)


def test_lo_que_alberto_marco_a_mano_y_la_norma(excel_caotico):
    assert excel_caotico.marcados_para_revisar() == {"PO-2026-0007", "PO-2026-0141"}
    assert excel_caotico.texto_norma()[1].startswith("1. Pagar solo")


def test_la_version_cambia_si_cambia_una_celda(excel_caotico, tmp_path):
    v1 = excel_caotico.version
    wb = openpyxl.load_workbook(excel_caotico.ruta)
    wb["Pedidos_2026"]["D2"] = 1.0
    wb.save(excel_caotico.ruta)
    assert MaestroExcel(excel_caotico.ruta).version != v1


@pytest.mark.skipif(not REAL.exists(), reason="La Caja no está clonada")
def test_el_excel_real_de_alberto():
    m = MaestroExcel(REAL)
    provs = {p.id: p for p in m.proveedores()}
    pedidos = {p.id: p for p in m.pedidos()}
    assert len(provs) == 11 and len(pedidos) == 518  # 516 de 2026 + 2 archivados de 2025
    assert provs["P003"].nombre == "Ofimática Cieza S.L."
    assert sum(1 for p in pedidos.values() if not p.nif and p.estado == "ABIERTO") == 20  # el bloque PO-0538..0557
    assert pedidos["PO-2026-0497"].importe == Decimal("84700.00")
    assert m.marcados_para_revisar() == {"PO-2026-0007", "PO-2026-0141"}
    assert len(m.texto_norma()) >= 7
    assert any("P007 repetido" in a for a in m.avisos)
